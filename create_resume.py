from docx import Document
from pathlib import Path
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 대상 MD 파일 목록
md_files = [
    "생성결과/생성_라이언로켓_중구.md",
    "생성결과/생성_RIPPLEAI_강남.md",
    "생성결과/생성_AITRICS_강남.md",
    "생성결과/생성_한국딥러닝_서초구.md",
    "생성결과/생성_하이퍼커넥트.md",
    "생성결과/생성_포지큐브_강남.md",
    "생성결과/생성_콕스웨이브_영등포.md",
    "생성결과/생성_인라이플_구로구.md",
    "생성결과/생성_이스트소프트_서초구.md",
    "생성결과/생성_머니투데이_종로.md",
    "생성결과/생성_메사쿠어컴퍼니_강남구.md",
    "생성결과/생성_메이아이_강남구.md",
    "생성결과/생성_버즈니_관악.md",
    "생성결과/생성_베슬AI_서초구.md",
    "생성결과/생성_aipm셀렉트스타_역삼.md",
    "생성결과/생성_셀렉트스타_역삼역.md",
    "생성결과/생성_모아데이타_성남 pm.md",
    "생성결과/생성_아하랩스_성남.md",
    "생성결과/생성_앨리스_선릉.md",
    "생성결과/생성_업스테이지_광교.md",
    "생성결과/생성_에임_강남.md",
    "생성결과/생성_지원동기_포지큐브_강남pm.md",
    
]

output_dir = "출력"
Path(output_dir).mkdir(parents=True, exist_ok=True)

def create_docx_from_md(md_filepath, output_dir):
    """MD 파일을 읽어 워드 문서(.docx)로 생성합니다."""
    with open(md_filepath, "r", encoding="utf-8") as f:
        md_content = f.read()

    document = Document()

    # 기본 스타일 설정
    normal_style = document.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = '맑은 고딕'
    normal_font.size = Pt(11)

    # 파일명에서 회사명 추출
    filename = Path(md_filepath).stem
    parts = filename.split("_")
    company_name = parts[1] if len(parts) >= 2 else "Unknown"

    document.add_heading(f"이력서_양호영_{company_name}", level=1)

    for line in md_content.splitlines():
        if line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("* "):
            p = document.add_paragraph(line[2:], style='List Bullet')
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT  # 왼쪽 정렬
        elif line.startswith("***"): # 구분선 처리
            document.add_paragraph()
        elif line.strip(): # 빈 줄이 아닌 경우 일반 텍스트로 추가
            p = document.add_paragraph(line, style='Normal') # 기본 스타일 적용
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT  # 왼쪽 정렬

    output_filename = f"이력서_양호영_{company_name}.docx"
    output_path = Path(output_dir) / output_filename
    document.save(output_path)
    print(f"📄 {output_path} 생성 완료")

# MD 파일 목록을 순회하며 DOCX 파일 생성
for md_file in md_files:
    create_docx_from_md(md_file, output_dir)

print("✅ 모든 이력서 생성 완료!")


