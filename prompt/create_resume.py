from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from pathlib import Path
import os
from docx2pdf import convert
import glob
import platform
import subprocess

# 폴더 구조 정의
INPUT_DIR = "C. 생성결과"        # 입력 파일이 있는 폴더
OUTPUT_BASE_DIR = "D. 출력"      # 기본 출력 폴더
OUTPUT_DOCX_DIR = os.path.join(OUTPUT_BASE_DIR, "docx")  # Word 문서 출력 폴더
OUTPUT_PDF_DIR = os.path.join(OUTPUT_BASE_DIR, "pdf")    # PDF 출력 폴더

# 출력 폴더들 생성
Path(OUTPUT_DOCX_DIR).mkdir(parents=True, exist_ok=True)
Path(OUTPUT_PDF_DIR).mkdir(parents=True, exist_ok=True)

def convert_to_pdf(docx_path):
    """Word 문서를 PDF로 변환합니다."""
    # PDF 파일 경로 생성
    pdf_path = str(docx_path).replace(OUTPUT_DOCX_DIR, OUTPUT_PDF_DIR).replace('.docx', '.pdf')
    
    print(f"\nPDF 변환 시작...")
    print(f"입력 파일: {docx_path}")
    print(f"출력 파일: {pdf_path}")
    
    # PDF 파일이 저장될 디렉토리 확인 및 생성
    pdf_dir = os.path.dirname(pdf_path)
    if not os.path.exists(pdf_dir):
        os.makedirs(pdf_dir)
        
    # Word 문서 존재 확인
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"Word 문서를 찾을 수 없습니다: {docx_path}")

    errors = []
    
    # 1. docx2pdf 시도
    try:
        convert(str(docx_path), str(pdf_path))
        if os.path.exists(pdf_path):
            print(f"✅ PDF 변환 완료 (docx2pdf 사용): {pdf_path}")
            return True
    except Exception as e:
        errors.append(f"docx2pdf 실패: {str(e)}")
        
    # 2. LibreOffice 시도
    try:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', pdf_dir, docx_path], 
                      capture_output=True, text=True, check=True)
        if os.path.exists(pdf_path):
            print(f"✅ PDF 변환 완료 (LibreOffice 사용): {pdf_path}")
            return True
    except Exception as e:
        errors.append(f"LibreOffice 실패: {str(e)}")
        
    # 3. unoconv 시도
    try:
        subprocess.run(['unoconv', '-f', 'pdf', '-o', pdf_path, docx_path], 
                      capture_output=True, text=True, check=True)
        if os.path.exists(pdf_path):
            print(f"✅ PDF 변환 완료 (unoconv 사용): {pdf_path}")
            return True
    except Exception as e:
        errors.append(f"unoconv 실패: {str(e)}")
        
    # 4. pandoc 시도
    try:
        subprocess.run(['pandoc', '-f', 'docx', '-t', 'pdf', '-o', pdf_path, docx_path],
                      capture_output=True, text=True, check=True)
        if os.path.exists(pdf_path):
            print(f"✅ PDF 변환 완료 (pandoc 사용): {pdf_path}")
            return True
    except Exception as e:
        errors.append(f"pandoc 실패: {str(e)}")

    # 모든 방법 실패
    error_msg = "\n".join(errors)
    print(f"⚠️ PDF 변환 실패. 다음 방법들을 시도했습니다:\n{error_msg}")
    
    print("\n💡 다음 명령어로 필요한 도구들을 설치해주세요:")
    if platform.system() == 'Darwin':  # macOS
        print("brew install unoconv libreoffice pandoc")
    elif platform.system() == 'Linux':
        print("sudo apt-get install unoconv libreoffice pandoc")
    else:  # Windows
        print("1. LibreOffice 설치: https://www.libreoffice.org/download/download/")
        print("2. Pandoc 설치: https://pandoc.org/installing.html")
    
    raise Exception("PDF 변환에 실패했습니다. 위 안내에 따라 필요한 도구를 설치해주세요.")

def ensure_conversion_tools():
    """필요한 변환 도구들을 확인하고 설치 안내를 제공합니다."""
    missing_tools = []
    
    # unoconv 확인
    if not subprocess.run(['where', 'unoconv'], capture_output=True).returncode == 0:
        missing_tools.append('unoconv')
    
    # libreoffice 확인
    if not subprocess.run(['where', 'soffice'], capture_output=True).returncode == 0:
        missing_tools.append('libreoffice')
    
    # pandoc 확인
    if not subprocess.run(['where', 'pandoc'], capture_output=True).returncode == 0:
        missing_tools.append('pandoc')
    
    if missing_tools:
        print("\n⚠️ 일부 PDF 변환 도구가 설치되어 있지 않습니다.")
        print("다음 도구를 설치하면 PDF 변환 성공률이 높아집니다:")
        
        if platform.system() == 'Darwin':  # macOS
            print("설치 명령어:")
            print("brew install " + " ".join(missing_tools))
        elif platform.system() == 'Linux':
            print("설치 명령어:")
            print("sudo apt-get install " + " ".join(missing_tools))
        else:  # Windows
            print("다음 도구를 설치해주세요:")
            for tool in missing_tools:
                print(f"- {tool}")

def create_word_from_md(md_file):
    """MD 파일을 읽어 Word 문서로 생성하고 PDF로 변환합니다."""
    md_filepath = os.path.join(INPUT_DIR, md_file)
    with open(md_filepath, "r", encoding="utf-8") as f:
        md_content = f.read()

    # 이름과 회사명 추출
    lines = md_content.splitlines()
    name = "Unknown"
    company_name = "Unknown"

    for line in lines:
        if line.startswith("**이름**:"):
            name = line.split(":")[1].strip()  # 이름 추출
        elif line.startswith("**회사명**:"):
            company_name = line.split(":")[1].strip()  # 회사명 추출

    output_filename = f"이력서_{name}_{company_name}.docx"  # 이름과 회사명을 포함한 파일명
    output_path = Path(OUTPUT_DOCX_DIR) / output_filename

    doc = Document()
    
    # 기본 폰트 및 크기 설정
    base_font_size = 12
    section_font_size = int(base_font_size * 1.5)  # 섹션 제목 크기
    name_font_size = 24
    
    # 문단 스타일 설정
    styles = doc.styles
    
    # 이름 스타일
    name_style = styles.add_style('CustomName', 1)
    name_font = name_style.font
    name_font.name = '맑은 고딕'
    name_font.size = Pt(name_font_size)
    name_font.bold = True
    
    # 섹션 제목 스타일
    section_style = styles.add_style('CustomSection', 1)
    section_font = section_style.font
    section_font.name = '맑은 고딕'
    section_font.size = Pt(section_font_size)
    section_font.bold = True
    
    # 일반 텍스트 스타일
    normal_style = styles.add_style('CustomNormal', 1)
    normal_font = normal_style.font
    normal_font.name = '맑은 고딕'
    normal_font.size = Pt(base_font_size)
    
    # 이름 추가
    p = doc.add_paragraph(style='CustomName')
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.add_run(name).bold = True
    doc.add_paragraph()  # 빈 줄 추가

    sections = ["핵심역량", "경력", "학력", "수상 및 기타", "병역"]
    first_line = True
    empty_line_count = 0
    
    lines = md_content.splitlines()
    for i, line in enumerate(lines):
        line = line.strip()
        
        # 빈 줄 처리
        if not line:
            empty_line_count += 1
            if empty_line_count == 1:
                doc.add_paragraph()  # 빈 줄 추가
            continue
        else:
            empty_line_count = 0
            
        if first_line:
            first_line = False
            continue
            
        if line.startswith("## "):
            # 섹션 제목 앞에 빈 줄 추가
            doc.add_paragraph()
            
            # 섹션 제목
            section_text = line[3:].replace(":", "").strip()
            if any(section in section_text for section in sections):
                p = doc.add_paragraph(style='CustomSection')
                p.add_run(section_text).bold = True
                doc.add_paragraph()  # 섹션 제목 뒤에 빈 줄 추가
            else:
                p = doc.add_paragraph(style='CustomNormal')
                p.add_run(section_text)
            
        elif line.startswith("* "):
            # 글머리 기호
            p = doc.add_paragraph(style='CustomNormal')
            p.paragraph_format.left_indent = Inches(0.25)
            text = line[2:]
            
            # 볼드 처리
            parts = text.split("**")
            if len(parts) >= 3:
                for i, part in enumerate(parts):
                    run = p.add_run(part)
                    run.bold = (i % 2 == 1)  # 홀수 인덱스의 부분을 볼드 처리
            else:
                p.add_run(text)
                
        elif line.startswith("***"):
            doc.add_paragraph()  # 빈 줄 추가
        else:
            # 일반 텍스트
            p = doc.add_paragraph(style='CustomNormal')
            
            # 볼드 처리
            parts = line.split("**")
            if len(parts) >= 3:
                for i, part in enumerate(parts):
                    run = p.add_run(part)
                    run.bold = (i % 2 == 1)  # 홀수 인덱스의 부분을 볼드 처리
            else:
                p.add_run(line)
            
            # 다음 줄이 새로운 문단이면 빈 줄 추가
            if i + 1 < len(lines):
                next_line = lines[i + 1].strip()
                if next_line and not next_line.startswith(("* ", "## ", "***")):
                    doc.add_paragraph()

    # 문서 저장
    doc.save(str(output_path))
    print(f"📄 Word 문서 생성 완료: {output_path}")
    
    # Word 문서를 PDF로 변환
    convert_to_pdf(output_path)

def process_existing_docx():
    """기존 Word 문서들을 PDF로 변환합니다."""
    print("\n기존 Word 문서 PDF 변환 시작...")
    docx_files = glob.glob(os.path.join(OUTPUT_DOCX_DIR, "*.docx"))
    
    if not docx_files:
        print("변환할 Word 문서가 없습니다.")
        return
        
    print(f"발견된 Word 문서: {len(docx_files)}개")
    for docx_file in docx_files:
        print(f"\n처리 중: {os.path.basename(docx_file)}")
        convert_to_pdf(docx_file)
    
    print("Word 문서 변환 완료")

def main():
    print(f"\n📂 작업 시작...")
    
    # 변환 도구 확인
    ensure_conversion_tools()
    
    # 입력 폴더 존재 확인
    if not os.path.exists(INPUT_DIR):
        print(f"⚠️ 입력 폴더가 없습니다: {INPUT_DIR}")
        print(f"💡 '{INPUT_DIR}' 폴더를 생성하고 변환할 마크다운(.md) 파일을 넣어주세요.")
        Path(INPUT_DIR).mkdir(parents=True, exist_ok=True)
        return
        
    # 처리할 MD 파일 확인
    md_files = [f for f in os.listdir(INPUT_DIR) if f.endswith('.md')]
    
    if not md_files:
        print(f"⚠️ 변환할 마크다운 파일이 없습니다.")
        if os.path.exists(OUTPUT_DOCX_DIR):
            process_existing_docx()
        return
    
    # MD 파일 처리
    print(f"\n📝 마크다운 파일 발견: {len(md_files)}개")
    print("\n변환 시작...")
    
    success_count = 0
    for md_file in md_files:
        try:
            print(f"\n처리 중: {md_file}")
            create_word_from_md(md_file)
            success_count += 1
        except Exception as e:
            print(f"⚠️ {md_file} 처리 중 오류 발생: {e}")
    
    if success_count == 0:
        print("\n❌ 모든 파일 변환에 실패했습니다.")
        raise Exception("PDF 변환에 실패했습니다. Microsoft Word가 설치되어 있는지 확인해주세요.")
    
    print(f"\n✅ {success_count}개 파일 변환 완료")

if __name__ == "__main__":
    main()
    print("\n✨ 모든 문서 변환 완료!")
    print(f"📁 출력 폴더를 확인해주세요:")
    print(f"  - Word 문서: {OUTPUT_DOCX_DIR}")
    print(f"  - PDF 파일: {OUTPUT_PDF_DIR}")
    print("\n💡 PDF 변환에 실패한 경우, Word 문서를 직접 열어서 PDF로 저장할 수 있습니다.")


