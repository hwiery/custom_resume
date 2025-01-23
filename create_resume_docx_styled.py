from docx import Document
from pathlib import Path
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

# 대상 MD 파일 목록 (기존과 동일)
md_files = [
    "생성결과/생성_라이언로켓_중구.md",
    "생성결과/생성_RIPPLEAI_강남.md",
    "생성결과/생성_AITRICS_강남.md",
    "생성결과/생성_한국딥러닝_서초구.md",
    "생성결과/생성_하이퍼커넥트.md",
    "생성결과/생성_포지큐브_강남.md",
    "생성결과/생성_콕스웨이브_영등포.md",
    "생성결과/생성_인라이플_구로구.md",
    "생성결과/생성_이스트소프트_서초구.md",
    "생성결과/생성_머니투데이_종로.md",
    "생성결과/생성_메사쿠어컴퍼니_강남구.md",
    "생성결과/생성_메이아이_강남구.md",
    "생성결과/생성_버즈니_관악.md",
    "생성결과/생성_베슬AI_서초구.md",
    "생성결과/생성_aipm셀렉트스타_역삼",
    "생성결과/생성_셀렉트스타_역삼역.md",
    "생성결과/생성_모아데이타_성남 pm.md",
]

output_dir = "출력"
Path(output_dir).mkdir(parents=True, exist_ok=True)

def create_docx_from_md_styled(md_filepath, output_dir):
    """MD 파일을 읽어 스타일이 적용된 워드 문서(.docx)로 생성합니다."""
    with open(md_filepath, "r", encoding="utf-8") as f:
        md_content = f.read()

    document = Document()

    # 색상 정의
    primary_color = RGBColor(0x33, 0x66, 0x99)

    # 기본 스타일 설정
    normal_style = document.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = '맑은 고딕'
    normal_font.size = Pt(11)

    # 제목 스타일 설정
    obj_styles = document.styles
    title_style = obj_styles.add_style('CustomTitle', WD_STYLE_TYPE.PARAGRAPH)
    title_font = title_style.font
    title_font.name = '맑은 고딕'
    title_font.size = Pt(18)
    title_font.bold = True
    title_font.color.rgb = primary_color

    # 부제목 스타일 설정
    subtitle_style = obj_styles.add_style('CustomSubtitle', WD_STYLE_TYPE.PARAGRAPH)
    subtitle_font = subtitle_style.font
    subtitle_font.name = '맑은 고딕'
    subtitle_font.size = Pt(14)
    subtitle_font.bold = True
    subtitle_font.color.rgb = primary_color

    # 목록 스타일 설정
    bullet_style = obj_styles.add_style('CustomListBullet', WD_STYLE_TYPE.PARAGRAPH)
    bullet_font = bullet_style.font
    bullet_font.name = '맑은 고딕'
    bullet_font.size = Pt(11)

    # 파일명에서 회사명 추출
    filename = Path(md_filepath).stem
    parts = filename.split("_")
    company_name = parts[1] if len(parts) >= 2 else "Unknown"

    p = document.add_paragraph(style='CustomTitle')
    p.add_run("이력서_양호영_").bold = True
    p.add_run(company_name).bold = True
    p.add_run("\n").bold = True

    for line in md_content.splitlines():
        if line.startswith("## "):
            document.add_paragraph(line[3:], style='CustomSubtitle')
        elif line.startswith("* "):
            p = document.add_paragraph(line[2:], style='CustomListBullet')
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith("***"):
            document.add_paragraph()
        elif line.strip():
            p = document.add_paragraph(line, style='Normal')
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    output_filename = f"이력서_양호영_{company_name}_styled.docx"
    output_path = Path(output_dir) / output_filename
    document.save(output_path)
    print(f"📄 {output_path} 생성 완료")

# MD 파일 목록을 순회하며 DOCX 파일 생성
for md_file in md_files:
    create_docx_from_md_styled(md_file, output_dir)

print("✅ 모든 이력서 생성 완료!") 